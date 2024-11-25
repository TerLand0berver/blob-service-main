"""
Document parser implementation.
"""
import io
from typing import Dict, Any, Optional
import magic
import docx
import PyPDF2
import csv
import openpyxl

from .base import BaseParser, ParserError

class DocumentParser(BaseParser):
    """Parser for document files (DOCX, PDF, CSV, XLSX, and text files)."""
    
    SUPPORTED_TYPES = {
        'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        'application/pdf',
        'text/csv',
        'text/plain'
    }
    
    SUPPORTED_EXTENSIONS = {
        '.docx', '.xlsx', '.pdf', '.csv', '.txt'
    }

    def __init__(self, config: Optional[Dict[str, Any]] = None):
        """Initialize document parser."""
        super().__init__(config)
        self.max_excel_rows = self.config.get('max_excel_rows', 1000)

    async def can_handle(self, content_type: str, file_extension: str) -> bool:
        """Check if parser can handle this file type."""
        return (content_type.lower() in self.SUPPORTED_TYPES or
                file_extension.lower() in self.SUPPORTED_EXTENSIONS)

    async def validate(self, file_data: bytes, metadata: Dict[str, Any]) -> bool:
        """Validate document format."""
        try:
            mime = magic.Magic(mime=True)
            content_type = mime.from_buffer(file_data)
            return content_type in self.SUPPORTED_TYPES
        except Exception:
            return False

    async def parse(self, file_data: bytes, metadata: Dict[str, Any]) -> Dict[str, Any]:
        """Parse document and extract text."""
        try:
            mime = magic.Magic(mime=True)
            content_type = mime.from_buffer(file_data)
            text_content = ""

            if content_type == 'application/vnd.openxmlformats-officedocument.wordprocessingml.document':
                doc = docx.Document(io.BytesIO(file_data))
                text_content = "\n".join([paragraph.text for paragraph in doc.paragraphs])
            
            elif content_type == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet':
                wb = openpyxl.load_workbook(io.BytesIO(file_data), read_only=True)
                texts = []
                for sheet in wb.sheetnames:
                    ws = wb[sheet]
                    sheet_texts = []
                    for idx, row in enumerate(ws.rows):
                        if idx >= self.max_excel_rows:
                            break
                        row_text = [str(cell.value) if cell.value is not None else '' for cell in row]
                        sheet_texts.append('\t'.join(row_text))
                    texts.append(f"Sheet: {sheet}\n" + '\n'.join(sheet_texts))
                text_content = '\n\n'.join(texts)
                wb.close()
            
            elif content_type == 'application/pdf':
                pdf = PyPDF2.PdfReader(io.BytesIO(file_data))
                text_content = "\n".join([page.extract_text() for page in pdf.pages])
            
            elif content_type == 'text/csv':
                decoded_content = file_data.decode('utf-8', errors='ignore')
                csv_reader = csv.reader(io.StringIO(decoded_content))
                text_content = "\n".join([",".join(row) for row in csv_reader])
            
            else:  # 普通文本文件
                text_content = file_data.decode('utf-8', errors='ignore')

            metadata.update({
                'content_type': content_type,
                'text_content': text_content,
                'size': len(file_data)
            })
            
            return metadata
        except Exception as e:
            raise ParserError(f"Failed to parse document: {e}")

    async def optimize(self, file_data: bytes, metadata: Dict[str, Any]) -> tuple[bytes, Dict[str, Any]]:
        """No optimization needed for documents."""
        return file_data, metadata
